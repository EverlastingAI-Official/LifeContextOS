"""Deterministic local document parsing for the HARNESS ingestion pipeline."""

from __future__ import annotations

import csv
import io
import json
import re
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path


SUPPORTED_EXTENSIONS = {".txt", ".md", ".json", ".csv", ".html", ".htm", ".pdf", ".docx"}


@dataclass(slots=True)
class TextUnit:
    locator: str
    text: str


class _HTMLTextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        value = data.strip()
        if value:
            self.parts.append(value)


def _read_text(path: Path) -> str:
    payload = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "utf-16"):
        try:
            return payload.decode(encoding)
        except UnicodeDecodeError:
            continue
    return payload.decode("utf-8", errors="replace")


def _flatten_json(value: object, prefix: str = "root") -> list[TextUnit]:
    units: list[TextUnit] = []
    if isinstance(value, dict):
        preferred = ["text", "content", "message", "title", "name"]
        text_parts = [str(value[key]) for key in preferred if isinstance(value.get(key), str)]
        if text_parts:
            units.append(TextUnit(prefix, "\n".join(text_parts)))
        for key, child in value.items():
            if key not in preferred:
                units.extend(_flatten_json(child, f"{prefix}.{key}"))
    elif isinstance(value, list):
        for index, child in enumerate(value):
            units.extend(_flatten_json(child, f"{prefix}[{index}]"))
    elif isinstance(value, str) and value.strip():
        units.append(TextUnit(prefix, value.strip()))
    return units


def parse_document(path: Path) -> list[TextUnit]:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {suffix or '(none)'}")
    if suffix in {".txt", ".md"}:
        return [TextUnit("document", _read_text(path))]
    if suffix == ".json":
        return _flatten_json(json.loads(_read_text(path)))
    if suffix == ".csv":
        rows = csv.reader(io.StringIO(_read_text(path)))
        return [TextUnit(f"row:{index}", " | ".join(row)) for index, row in enumerate(rows, 1)]
    if suffix in {".html", ".htm"}:
        parser = _HTMLTextExtractor()
        parser.feed(_read_text(path))
        return [TextUnit("document", "\n".join(parser.parts))]
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return [TextUnit(f"page:{index}", page.extract_text() or "") for index, page in enumerate(reader.pages, 1)]
    if suffix == ".docx":
        from docx import Document

        document = Document(str(path))
        units = [TextUnit(f"paragraph:{index}", paragraph.text) for index, paragraph in enumerate(document.paragraphs, 1) if paragraph.text.strip()]
        for table_index, table in enumerate(document.tables, 1):
            for row_index, row in enumerate(table.rows, 1):
                text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if text:
                    units.append(TextUnit(f"table:{table_index}/row:{row_index}", text))
        return units
    raise AssertionError("unreachable")


def chunk_units(units: list[TextUnit], target_chars: int = 1800) -> list[TextUnit]:
    chunks: list[TextUnit] = []
    for unit in units:
        clean = re.sub(r"\n{3,}", "\n\n", unit.text).strip()
        if not clean:
            continue
        paragraphs = re.split(r"(?<=\n)|(?<=[。！？.!?])", clean)
        buffer = ""
        part = 1
        for paragraph in paragraphs:
            if buffer and len(buffer) + len(paragraph) > target_chars:
                chunks.append(TextUnit(f"{unit.locator}/chunk:{part}", buffer.strip()))
                part += 1
                buffer = ""
            buffer += paragraph
        if buffer.strip():
            chunks.append(TextUnit(f"{unit.locator}/chunk:{part}", buffer.strip()))
    return chunks
